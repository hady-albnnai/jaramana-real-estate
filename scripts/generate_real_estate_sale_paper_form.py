from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/forms/real_estate_sale_intake_sweida_print.docx'
LOGO = ROOT / 'assets/images/logo.png'
GOLD = 'B8892F'
INK = '171717'
SOFT = 'F8EFD9'
LINE = 'D8C69C'


def rtl(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = p._p.get_or_add_pPr()
    bidi = pPr.find(qn('w:bidi'))
    if bidi is None:
        bidi = OxmlElement('w:bidi')
        pPr.append(bidi)
    return p


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color=LINE, size='6'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        elem = borders.find(qn(tag))
        if elem is None:
            elem = OxmlElement(tag)
            borders.append(elem)
        elem.set(qn('w:val'), 'single')
        elem.set(qn('w:sz'), size)
        elem.set(qn('w:color'), color)


def setup_doc():
    d = Document()
    sec = d.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(0.8)
    sec.bottom_margin = Cm(1.2)
    sec.right_margin = Cm(1.1)
    sec.left_margin = Cm(1.1)
    styles = d.styles
    normal = styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(9)
    normal._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0
    return d


def header(d, title, subtitle='', form_copy=''):
    t = d.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Cm(3.0)
    t.columns[1].width = Cm(12.4)
    t.columns[2].width = Cm(3.2)
    c0, c1, c2 = t.rows[0].cells
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Cm(2.4))
    p = rtl(c1.paragraphs[0])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(GOLD)
    if subtitle:
        p2 = rtl(c1.add_paragraph(subtitle))
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.runs[0].font.size = Pt(8)
    p = rtl(c2.paragraphs[0])
    p.add_run('رقم الاستمارة\n').bold = True
    p.add_run(form_copy or '........................\n')
    p.add_run('التاريخ: .... / .... / 2026')
    for c in (c0, c1, c2):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(c, GOLD, '8')
    d.add_paragraph()


def section_title(d, n, text):
    t = d.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    set_cell_shading(c, SOFT)
    set_cell_border(c, GOLD, '8')
    p = rtl(c.paragraphs[0])
    r = p.add_run(f'{n}) {text}')
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(INK)


def fields(d, items, cols=2):
    rows = (len(items) + cols - 1) // cols
    t = d.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i in range(rows * cols):
        c = t.cell(i // cols, i % cols)
        set_cell_border(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = rtl(c.paragraphs[0])
        if i < len(items):
            label, lines = items[i] if isinstance(items[i], tuple) else (items[i], 1)
            r = p.add_run(label)
            r.bold = True
            r.font.size = Pt(8.5)
            for _ in range(lines):
                q = rtl(c.add_paragraph('........................................................................................'))
                q.paragraph_format.space_after = Pt(0)
                q.runs[0].font.size = Pt(8)
        else:
            p.add_run('')
    d.add_paragraph()


def options(d, title, values):
    p = rtl(d.add_paragraph())
    r = p.add_run(title + ':  ')
    r.bold = True
    r.font.size = Pt(8.5)
    p.add_run('    '.join('☐ ' + v for v in values))


def note(d, text, warning=False):
    t = d.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    set_cell_shading(c, 'FFF4F2' if warning else 'FAF7EF')
    set_cell_border(c, '8A2E2E' if warning else GOLD, '8')
    p = rtl(c.paragraphs[0])
    r = p.add_run(text)
    r.font.size = Pt(8)
    if warning:
        r.font.color.rgb = RGBColor(138, 46, 46)


def checklist(d, headers, rows, widths=None):
    t = d.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        set_cell_shading(c, SOFT)
        p = rtl(c.paragraphs[0]); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(h); rr.bold = True; rr.font.size = Pt(8)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            p = rtl(cells[i].paragraphs[0])
            if i > 0 and v in ('☐', 'نعم / لا'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(v).font.size = Pt(8)
    d.add_paragraph()


def page_break(d):
    d.add_page_break()


def add_footer(section):
    f = section.footer
    p = rtl(f.paragraphs[0])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('المكتب العقاري الإلكتروني — السويداء   |   0958750199   |   صفحة ')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    p._p.append(fld)


def build():
    d = setup_doc()
    add_footer(d.sections[0])
    header(d, 'استمارة عرض بيع عقار', 'نسخة ورقية عامة — متوافقة مع بيانات التطبيق وواقع السويداء')
    note(d, 'الحقول الأساسية إلزامية لإدخال العرض. بيانات الهوية والسند والموقع الدقيق خاصة بالمكتب ولا تُنشر للعموم.')
    section_title(d, 1, 'معلومات المالك أو مقدم الطلب')
    fields(d, [('الاسم الكامل كما في الهوية *', 1), ('رقم الهاتف الرئيسي *', 1), ('رقم هاتف بديل', 1), ('الأوقات المناسبة للتواصل', 1), ('الرقم الوطني / رقم الهوية', 1), ('عنوان الإقامة', 1)], 2)
    options(d, 'الصفة', ['المالك المسجل', 'وكيل قانوني', 'أحد الورثة', 'شريك / مالك على الشيوع', 'أخرى: ................'])
    fields(d, [('اسم المالك المسجل إن كان مقدم الطلب وكيلاً', 1), ('صلة الوكيل بالمالك وبيانات الوكالة', 1)], 2)
    section_title(d, 2, 'تعريف العرض والتصنيف')
    fields(d, [('عنوان مختصر للعرض داخل التطبيق *', 1), ('موعد الإخلاء والتسليم', 1)], 2)
    options(d, 'سكني', ['شقة سكنية', 'دار عربي', 'فيلا', 'مزرعة', 'بناء كامل', 'سطح'])
    options(d, 'تجاري', ['محل', 'معرض', 'مركز تجاري', 'مكتب', 'مستودع'])
    options(d, 'زراعي', ['أرض زراعية', 'دواجن', 'مواشي', 'مشتل'])
    options(d, 'صناعي', ['منشأة', 'ورشة', 'مصنع', 'أرض صناعية', 'آخر: ................'])
    options(d, 'حالة الإشغال', ['فارغ', 'يسكنه المالك', 'مؤجر', 'أخرى'])
    section_title(d, 3, 'السعر وشروط البيع')
    fields(d, [('السعر المطلوب إجمالاً *', 1), ('أقل سعر مقبول — سري', 1), ('تاريخ تثبيت السعر', 1), ('تفاصيل الأقساط أو المقايضة', 2)], 2)
    options(d, 'العملة', ['دولار أمريكي', 'ليرة سورية'])
    options(d, 'الشروط', ['قابل للتفاوض', 'نقداً بالكامل', 'أقساط', 'مقايضة بعقار', 'محسوب على المتر'])
    fields(d, [('اتفاق أتعاب / عمولة المكتب — داخلي', 2)], 1)

    page_break(d)
    header(d, 'موقع العقار وبيانات البناء', 'الموقع العام للنشر — والتفصيل الدقيق للتنسيق فقط')
    section_title(d, 4, 'الموقع')
    fields(d, [('المحافظة *', 1), ('المدينة / المنطقة الرئيسية *', 1), ('البلدة / القرية', 1), ('الحي / المنطقة العقارية *', 1), ('اسم الشارع', 1), ('أقرب نقطة دالة', 1), ('عرض الطريق وحالته', 1), ('رابط أو إحداثيات الموقع — خاص', 1)], 2)
    fields(d, [('الوصف الدقيق للوصول — خاص بالمكتب والمصور', 3)], 1)
    section_title(d, 5, 'المساحات والبناء والتوزيع')
    fields(d, [('مساحة العقار / الأرض م² *', 1), ('المساحة المبنية الإجمالية م²', 1), ('المساحة الصافية م²', 1), ('طول الواجهة / عرضها', 1), ('عمر البناء', 1), ('عدد طوابق البناء', 1), ('رقم الطابق', 1), ('عدد المداخل', 1), ('عدد الشرفات', 1), ('الارتفاع الصافي', 1), ('غرف النوم', 1), ('الصالونات', 1), ('غرف المعيشة', 1), ('المطابخ', 1), ('الحمامات', 1), ('دورات المياه', 1)], 4)
    options(d, 'الاتجاهات', ['قبلي / جنوبي', 'شمالي', 'شرقي', 'غربي', 'إطلالة مفتوحة', 'زاوية'])
    options(d, 'الإكساء', ['ملكي', 'سوبر ديلوكس', 'ديلوكس', 'عادي', 'هيكل / عظم', 'قديم يحتاج ترميم'])
    fields(d, [('وصف التوزيع الداخلي', 3), ('الترميمات وحالتها وتاريخها', 3)], 2)

    page_break(d)
    header(d, 'التجهيزات والخدمات', 'الكهرباء والطاقة والمياه والتدفئة والخدمات المشتركة')
    section_title(d, 6, 'الخدمات الأساسية')
    checklist(d, ['الخدمة', 'متوفر', 'التفاصيل'], [
        ('كهرباء نظامية — عداد مستقل / مشترك / ثلاثي الطور', '☐', '........................................................'),
        ('طاقة شمسية — ألواح / إنفرتر / بطاريات', '☐', '........................................................'),
        ('مولدة أو مصدر بديل', '☐', '........................................................'),
        ('مياه الشبكة وعداد مستقل ودورية الضخ', '☐', '........................................................'),
        ('خزان أرضي / علوي — العدد والسعة والعزل', '☐', '........................................................'),
        ('بئر — الترخيص والغزارة والعمق والمضخة', '☐', '........................................................'),
        ('تدفئة مركزية / مازوت / كهرباء / وجاق', '☐', '........................................................'),
        ('صرف صحي نظامي', '☐', '........................................................'),
        ('إنترنت وهاتف أرضي', '☐', '........................................................'),
        ('مصعد وحالته', '☐', '........................................................'),
    ])
    section_title(d, 7, 'مزايا السكني')
    options(d, 'المزايا', ['كراج خاص', 'موقف مفتوح', 'حديقة', 'قبو / مستودع', 'كاميرات', 'مفروش', 'مطبخ راكب', 'واجهة حجر', 'عزل', 'مسبح'])
    fields(d, [('ملكية الكراج ووصفه', 1), ('رسوم خدمات البناء', 1), ('التجهيزات التي تبقى ضمن البيع', 2), ('العيوب المعروفة أو الأعمال المطلوبة', 2)], 2)

    page_break(d)
    header(d, 'الوضع القانوني والملكية', 'لا يُعتمد وصف السند من دون الاطلاع على الوثائق')
    note(d, 'يجب استخراج بيان قيد حديث والتحقق من المالك والإشارات والرهون والدعاوى والحصص قبل أي عربون أو التزام.', True)
    section_title(d, 8, 'نوع السند — مطابق للتطبيق')
    options(d, 'السند', ['طابو أخضر', 'حصة سهمية — حكم محكمة', 'حصة سهمية — كاتب بالعدل', 'مستملك', 'تسلسل عقود', 'جمعية سكنية', 'نمرة قديمة', 'نمرة جديدة', 'وارد'])
    fields(d, [('المنطقة العقارية *', 1), ('رقم العقار / المحضر *', 1), ('رقم المقسم / الشقة', 1), ('الأسهم المملوكة من أصل', 1), ('اسم المالك في السجل', 1), ('تاريخ الوثيقة', 1)], 2)
    options(d, 'الملكية', ['مفرز ومستقل', 'على الشيوع', 'موافقة جميع الشركاء', 'انتقال الإرث مكتمل', 'الوكالة سارية'])
    section_title(d, 9, 'الإشارات والترخيص')
    options(d, 'الإشارات', ['لا يوجد حسب بيان حديث', 'رهن', 'حجز', 'إشارة دعوى', 'إشارة بيع', 'حق ارتفاق', 'غير معلوم'])
    fields(d, [('تفاصيل الإشارات وكيفية فكها', 3), ('الذمم المالية والرسوم المتبقية', 3)], 2)
    options(d, 'التنظيم', ['داخل التنظيم', 'خارج التنظيم', 'البناء مرخص', 'غير مرخص', 'مفرز', 'رخصة إشغال', 'توجد مخالفة'])
    fields(d, [('وصف المخالفة أو التسوية', 2)], 1)
    section_title(d, 10, 'الوثائق المرفقة')
    checklist(d, ['الوثيقة', 'مرفقة', 'مطابقة', 'ملاحظة'], [
        ('هوية المالك / الوكيل', '☐', '☐', '................................'),
        ('السند / الحكم / الوكالة كاملة', '☐', '☐', '................................'),
        ('بيان قيد حديث', '☐', '☐', '................................'),
        ('مخطط مساحة / إفراز / ترخيص', '☐', '☐', '................................'),
        ('حصر إرث وموافقات', '☐', '☐', '................................'),
        ('براءة ذمة / فواتير', '☐', '☐', '................................'),
    ])

    page_break(d)
    header(d, 'حقول الأرض والتجاري والوصف', 'يُملأ القسم المناسب فقط')
    section_title(d, 11, 'أرض زراعية / مزرعة')
    options(d, 'الاستخدام', ['زراعي', 'سكني', 'صناعي', 'تجاري', 'متعدد', 'مسيجة', 'مستوية'])
    fields(d, [('واجهة الأرض على الطريق', 1), ('نوع الطريق', 1), ('طبيعة التربة والميلان', 1), ('نوع وعدد الأشجار', 1), ('مصدر الري وحقوق المياه', 1), ('منشآت قائمة', 1), ('الخدمات القريبة وإمكانية إيصالها', 2), ('الاستعمال الحالي والجدوى المقترحة', 2)], 2)
    section_title(d, 12, 'تجاري / صناعي / مستودع')
    fields(d, [('طول الواجهة التجارية', 1), ('عدد الواجهات / الأبواب', 1), ('الارتفاع الصافي', 1), ('حمولة الأرضية', 1)], 2)
    options(d, 'التجهيزات', ['ترخيص مهني', 'كهرباء صناعية', 'مدخل تحميل', 'وصول شاحنات', 'دورة مياه', 'مكتب إداري', 'مؤجر'])
    fields(d, [('النشاط المرخص والمسموح', 2), ('تفاصيل الإيجار الحالي والعائد', 2)], 2)
    section_title(d, 13, 'الوصف التسويقي الكامل')
    fields(d, [('وصف العرض كما سيظهر في التطبيق *', 7), ('ملاحظات قانونية أو خاصة للمكتب — لا تُنشر تلقائياً', 3)], 1)

    page_break(d)
    header(d, 'التصوير والمعاينة والنشر', 'تحضير عرض كامل قابل للمراجعة والنشر')
    section_title(d, 14, 'خطة التصوير')
    options(d, 'المصدر', ['صور المالك', 'مصور المكتب', 'كلاهما', 'فيديو مطلوب'])
    checklist(d, ['اللقطة المطلوبة', 'تمت', 'ملاحظة'], [
        ('واجهة البناء والشارع', '☐', '........................................'),
        ('المدخل والدرج والمصعد', '☐', '........................................'),
        ('الغرف والصالونات', '☐', '........................................'),
        ('المطبخ والحمامات', '☐', '........................................'),
        ('الشرفات والإطلالة والكراج', '☐', '........................................'),
        ('الطاقة والخزانات والتدفئة', '☐', '........................................'),
        ('للأرض: الحدود والطريق والأشجار والبئر', '☐', '........................................'),
    ])
    section_title(d, 15, 'مواعيد المعاينة')
    options(d, 'الجاهزية', ['بأي وقت بعد التنسيق', 'أيام وفترات محددة'])
    checklist(d, ['اليوم', 'من', 'إلى'], [(x, '........', '........') for x in ['الاثنين','الثلاثاء','الأربعاء','الخميس','الجمعة','السبت','الأحد']])
    fields(d, [('الشخص الذي يفتح العقار للمعاينة', 1), ('رقمه وعلاقته بالمالك', 1)], 2)
    section_title(d, 16, 'موافقات النشر')
    options(d, 'الموافقات', ['النشر في التطبيق', 'النشر في صفحات المكتب', 'دخول المصور', 'وضع علامة المكتب'])
    fields(d, [('معلومات يجب عدم نشرها', 2), ('نقاط يرغب المالك بإبرازها', 2)], 2)

    page_break(d)
    header(d, 'الإقرارات والتوقيعات', 'تُقرأ البنود مع صاحب العرض قبل التوقيع')
    section_title(d, 17, 'إقرار صاحب العرض')
    for text in [
        'أقر بصحة جميع البيانات والمواصفات والأسعار والمعلومات القانونية المقدمة وتحت مسؤوليتي.',
        'أفصحت عن الملكية والحصص والورثة والإشارات والرهون والدعاوى والمخالفات والعيوب المعروفة.',
        'أفوض المكتب بالتحقق من الوثائق وترتيب التصوير والمعاينات، ولا يعني النشر ضمان سلامة الملكية قبل التدقيق.',
        'أوافق على نشر البيانات العامة والصور، مع بقاء الهوية والهاتف والموقع الدقيق والوثائق للاستخدام الداخلي.',
        'ألتزم بإبلاغ المكتب فور تغير السعر أو توفر العقار أو بيعه أو ظهور مانع قانوني أو مادي.',
    ]:
        p = rtl(d.add_paragraph('☐ ' + text)); p.runs[0].font.size = Pt(8.5)
    section_title(d, 18, 'التوقيع')
    fields(d, [('الاسم الثلاثي *', 1), ('الصفة', 1), ('التاريخ', 1), ('التوقيع / البصمة', 3)], 2)
    section_title(d, 19, 'مراجعة المكتب')
    checklist(d, ['البند', 'مكتمل', 'الملاحظة'], [
        ('بيانات الاتصال والصفة', '☐', '................................................'),
        ('التصنيف والسعر والعملة', '☐', '................................................'),
        ('الموقع والمساحة والمواصفات', '☐', '................................................'),
        ('السند والوثائق والإشارات', '☐', '................................................'),
        ('الصور والوصف ومواعيد المعاينة', '☐', '................................................'),
    ])
    options(d, 'القرار', ['جاهز للإدخال', 'نواقص', 'تدقيق قانوني', 'يحتاج تصويراً', 'غير مقبول'])
    fields(d, [('رقم العرض بعد الإدخال', 1), ('النواقص وخطة المتابعة', 3), ('اسم وتوقيع الموظف', 2), ('اسم وتوقيع المدقق', 2)], 2)
    note(d, 'هذه الاستمارة لجمع البيانات وتحضير الإعلان، وليست عقد بيع ولا سند ملكية ولا بديلاً عن بيان القيد أو الاستشارة القانونية والفنية.')
    d.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
