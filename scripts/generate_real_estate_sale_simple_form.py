from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/forms/real_estate_sale_public_simple.docx'
LOGO = ROOT / 'assets/images/logo.png'
GOLD = 'B8892F'; SOFT = 'F8EFD9'; LINE = 'D8C69C'


def rtl(p, center=False):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT
    pPr = p._p.get_or_add_pPr(); bidi = OxmlElement('w:bidi'); pPr.append(bidi)
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 0.9
    return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)


def border(cell, color=LINE):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for e in ('top','left','bottom','right'):
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'6'); x.set(qn('w:color'),color); b.append(x)


def title(doc, text):
    t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,SOFT); border(c,GOLD)
    p=rtl(c.paragraphs[0]); r=p.add_run(text); r.bold=True; r.font.size=Pt(9.5)


def lines(doc, labels, cols=2):
    rows=(len(labels)+cols-1)//cols; t=doc.add_table(rows=rows,cols=cols); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i in range(rows*cols):
        c=t.cell(i//cols,i%cols); border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        p=rtl(c.paragraphs[0])
        if i<len(labels):
            r=p.add_run(labels[i]+'\n'); r.bold=True; r.font.size=Pt(7.5)
            q=rtl(c.add_paragraph('................................................................................')); q.runs[0].font.size=Pt(7)


def options(doc, heading, values):
    p=rtl(doc.add_paragraph()); r=p.add_run(heading+':  '); r.bold=True; r.font.size=Pt(7.5)
    x=p.add_run('   '.join('☐ '+v for v in values)); x.font.size=Pt(7.5)


def build():
    d=Document(); s=d.sections[0]
    s.page_width=Cm(21); s.page_height=Cm(29.7); s.top_margin=Cm(.6); s.bottom_margin=Cm(.8); s.right_margin=Cm(.7); s.left_margin=Cm(.7)
    n=d.styles['Normal']; n.font.name='Arial'; n.font.size=Pt(7.5); n._element.rPr.rFonts.set(qn('w:cs'),'Arial'); n.paragraph_format.space_after=Pt(0)
    h=d.add_table(rows=1,cols=3); h.alignment=WD_TABLE_ALIGNMENT.CENTER
    c0,c1,c2=h.rows[0].cells
    p=rtl(c0.paragraphs[0],True); p.add_run().add_picture(str(LOGO),width=Cm(1.9))
    p=rtl(c1.paragraphs[0],True); r=p.add_run('استمارة بيانات عقار للبيع'); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string(GOLD)
    q=rtl(c1.add_paragraph('ورقة مبسطة لجمع المعلومات الأولية — المكتب العقاري الالكتروني - جرمانا'),True); q.runs[0].font.size=Pt(7)
    p=rtl(c2.paragraphs[0]); p.add_run('رقم: ................\nالتاريخ: .... / .... / 2026').font.size=Pt(7)
    for c in (c0,c1,c2): border(c,GOLD); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

    title(d,'1) معلومات التواصل'); lines(d,['الاسم','رقم الهاتف الأساسي *','رقم بديل','الوقت المناسب للاتصال'],2)
    title(d,'2) نوع العقار وموقعه')
    options(d,'النوع',['شقة','دار عربي','فيلا','أرض','مزرعة','محل','مكتب','مستودع','بناء كامل','أخرى'])
    lines(d,['المدينة / البلدة *','الحي / المنطقة *','الشارع أو أقرب نقطة دالة'],3)
    title(d,'3) معلومات العقار الأساسية')
    lines(d,['المساحة م² *','رقم الطابق','عدد الغرف','عدد الحمامات'],4)
    options(d,'الإكساء',['ملكي','سوبر ديلوكس','ديلوكس','عادي','هيكل / عظم','يحتاج ترميم'])
    options(d,'الاتجاه',['قبلي','شمالي','شرقي','غربي'])
    options(d,'الحالة',['فارغ','يسكنه المالك','مؤجر'])
    title(d,'4) الملكية والسعر')
    options(d,'السند',['طابو أخضر','حكم محكمة','حصة سهمية','كاتب بالعدل','تسلسل عقود','جمعية','أخرى'])
    lines(d,['رقم العقار / المحضر','السعر المطلوب *','العملة'],3)
    options(d,'الشروط',['قابل للتفاوض','نقداً','أقساط','مقايضة','توجد إشارة / رهن / مشكلة قانونية'])
    title(d,'5) الخدمات والميزات الموجودة')
    options(d,'الخدمات',['كهرباء مستقلة','طاقة شمسية','مياه شبكة','خزان أرضي','خزان علوي','بئر','تدفئة','مصعد','كراج','حديقة','مفروش','مطبخ راكب'])
    lines(d,['تفاصيل مهمة عن الماء أو الطاقة أو التدفئة أو الكراج'],1)
    title(d,'6) وصف مختصر وملاحظات')
    t=d.add_table(rows=1,cols=1); c=t.cell(0,0); border(c); p=rtl(c.paragraphs[0]); p.add_run('\n........................................................................................................................................................................\n........................................................................................................................................................................\n........................................................................................................................................................................').font.size=Pt(7)
    title(d,'7) الصور والمعاينة')
    options(d,'الصور',['لدي صور جاهزة','أحتاج مصور المكتب'])
    options(d,'المعاينة',['بأي وقت بعد التنسيق','تحتاج موعداً مسبقاً'])
    lines(d,['الأيام أو الأوقات المناسبة للمعاينة'],1)
    p=rtl(d.add_paragraph()); r=p.add_run('☐ أقر بأن البيانات المكتوبة صحيحة بحسب علمي، وأوافق على تواصل المكتب معي للتحقق والتصوير وتجهيز الإعلان. هذه الورقة لجمع بيانات أولية وليست عقد بيع أو سند ملكية.'); r.font.size=Pt(6.8)
    lines(d,['الاسم والتوقيع','ملاحظات موظف المكتب'],2)
    f=s.footer; p=rtl(f.paragraphs[0],True); r=p.add_run('المكتب العقاري الالكتروني - جرمانا   |   [رقم المكتب الرسمي]   |   بياناتك الخاصة لا تُنشر للعموم'); r.font.size=Pt(7); r.font.color.rgb=RGBColor.from_string(GOLD)
    d.save(OUT); print(OUT)

if __name__=='__main__': build()
